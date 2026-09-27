"""Lectura del Excel del caso (data/raw) hacia DataFrames crudos (data/interim).

Las hojas del Excel tienen un formato "de presentación": columna A vacía, 2-4 filas de
título y la cabecera más abajo. Aquí solo se localiza la tabla; no se cambian valores.
La limpieza semántica (tipos, nombres, validaciones) vive en `optimum.cleaning`.
"""

from __future__ import annotations

from pathlib import Path

import openpyxl
import pandas as pd

from optimum import paths


def _find_header_row(rows: list[tuple]) -> int:
    """Primera fila con >= 2 celdas no vacías: las filas de título solo tienen una."""
    for i, row in enumerate(rows):
        if sum(v is not None for v in row) >= 2:
            return i
    raise ValueError("No se encontró una fila de cabecera")


def read_sheet(ws) -> pd.DataFrame:
    rows = list(ws.iter_rows(values_only=True))
    h = _find_header_row(rows)
    header = rows[h]
    keep = [j for j, v in enumerate(header) if v is not None]
    cols = [str(header[j]).strip() for j in keep]
    body = [[r[j] for j in keep] for r in rows[h + 1 :]]
    df = pd.DataFrame(body, columns=cols)
    return df.dropna(how="all").reset_index(drop=True)


def read_sheet_notes(ws) -> list[str]:
    """Textos de título/nota sobre la cabecera (útiles como metadatos)."""
    rows = list(ws.iter_rows(values_only=True))
    h = _find_header_row(rows)
    return [v for r in rows[:h] for v in r if isinstance(v, str)]


def load_workbook(path: Path = paths.RAW_EXCEL) -> dict[str, pd.DataFrame]:
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    return {ws.title: read_sheet(ws) for ws in wb.worksheets}


def sheet_slug(name: str) -> str:
    return name.lower().replace("é", "e").replace(" ", "_").replace("-", "_")


def export_interim(path: Path = paths.RAW_EXCEL, out_dir: Path = paths.INTERIM) -> list[Path]:
    """Guarda cada hoja como parquet en data/interim/. Idempotente."""
    out_dir.mkdir(parents=True, exist_ok=True)
    written = []
    for name, df in load_workbook(path).items():
        # parquet exige columnas homogéneas: las columnas mixtas se guardan como texto
        df = df.copy()
        for c in df.columns:
            if df[c].dtype == object and df[c].map(type).nunique() > 1:
                df[c] = df[c].astype("string")
        f = out_dir / f"{sheet_slug(name)}.parquet"
        df.to_parquet(f, index=False)
        written.append(f)
    return written
